"""Tests de l'import multi-documents par lot dans DocumentsView.

Couvre la sélection multiple du dialogue natif, le glisser-déposer de fichiers
(vue + arbre) et l'orchestration séquentielle résiliente du lot.
"""

import pathlib
import threading
from unittest.mock import patch

import pytest
from PySide6.QtCore import QMimeData, QPoint, QPointF, Qt, QUrl
from PySide6.QtGui import QDragEnterEvent, QDropEvent
from PySide6.QtWidgets import QTreeWidgetItem

from ankiforge.database.models import DocumentModel
from ankiforge.services.workers.document_batch_worker import (
    DocumentBatchTask,
    DocumentBatchWorker,
)
from ankiforge.ui.views.documents_view import DocumentsView
from ankiforge.ui.views.documents_view.view import batch_summary_message
from ankiforge.ui.widgets.toast import ToastLevel, ToastManager

pytestmark = pytest.mark.ui

DIALOG_PATH = "ankiforge.ui.views.documents_view.view.QFileDialog"
WORKER_PATH = "ankiforge.ui.views.documents_view.view.DocumentBatchWorker"


def _make_files(tmp_path, names: list[str]) -> list[str]:
    paths = []
    for name in names:
        target = tmp_path / name
        target.write_text(f"# {name}\n\nContenu de {name}.", encoding="utf-8")
        paths.append(str(target))
    return paths


def _mime(paths: list[str]) -> QMimeData:
    mime = QMimeData()
    mime.setUrls([QUrl.fromLocalFile(p) for p in paths])
    return mime


def _drop_event(paths: list[str]) -> tuple[QDropEvent, QMimeData]:
    """Construit un QDropEvent ; le QMimeData doit rester vivant côté Python (pointeur C++)."""
    mime = _mime(paths)
    event = QDropEvent(
        QPointF(10, 10),
        Qt.DropAction.CopyAction,
        mime,
        Qt.MouseButton.LeftButton,
        Qt.KeyboardModifier.NoModifier,
    )
    return event, mime


def _drag_enter_event(paths: list[str]) -> tuple[QDragEnterEvent, QMimeData]:
    mime = _mime(paths)
    event = QDragEnterEvent(
        QPoint(10, 10),
        Qt.DropAction.CopyAction,
        mime,
        Qt.MouseButton.LeftButton,
        Qt.KeyboardModifier.NoModifier,
    )
    return event, mime


def _attach_batch_worker(view: DocumentsView, paths: list[str]) -> DocumentBatchWorker:
    """Rattache un worker de lot réel (jamais démarré) à la vue pour piloter ses signaux."""
    tasks = [DocumentBatchTask(index=i, path=p) for i, p in enumerate(paths)]
    worker = DocumentBatchWorker(tasks)
    view._bind_batch_worker(worker, tasks)
    return worker


# ── Sélection multiple via le dialogue natif ────────────────────────────────


def test_import_file_dialog_allows_multiple_selection(qtbot, tmp_path) -> None:
    """Le dialogue natif est ouvert en mode multi-sélection et alimente la file par lot."""
    paths = _make_files(tmp_path, ["un.md", "deux.md", "trois.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    with (
        patch(DIALOG_PATH) as dialog,
        patch.object(view, "_start_batch_import") as start,
    ):
        dialog.getOpenFileNames.return_value = (paths, "")
        view._on_import_file()

    assert dialog.getOpenFileName.called is False
    start.assert_called_once_with(paths)


def test_import_file_keeps_single_file_flow(qtbot, tmp_path) -> None:
    """Une sélection unitaire conserve le parcours historique (viewer PDF inclus)."""
    paths = _make_files(tmp_path, ["unique.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    started: list[str] = []
    with (
        patch(DIALOG_PATH) as dialog,
        patch.object(view, "_start_document_worker", side_effect=lambda p, doc_id=None: started.append(p)),
        patch.object(view, "_import_pdf_directly", side_effect=lambda p: started.append(p)),
        patch.object(view, "_start_batch_import") as start,
    ):
        dialog.getOpenFileNames.return_value = (paths, "")
        view._on_import_file()

    assert started == paths
    assert start.called is False


def test_import_file_without_selection_is_a_noop(qtbot) -> None:
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    with patch(DIALOG_PATH) as dialog, patch.object(view, "_start_batch_import") as start:
        dialog.getOpenFileNames.return_value = ([], "")
        view._on_import_file()

    assert start.called is False


# ── Glisser-déposer ─────────────────────────────────────────────────────────


def test_view_accepts_file_drag_and_drops_every_file(qtbot, tmp_path) -> None:
    paths = _make_files(tmp_path, ["drag1.md", "drag2.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    with patch.object(view, "_start_batch_import") as start:
        enter, _enter_mime = _drag_enter_event(paths)
        view.dragEnterEvent(enter)
        assert enter.isAccepted()
        drop, _drop_mime = _drop_event(paths)
        view.dropEvent(drop)

    start.assert_called_once_with(paths)


def test_view_rejects_drag_without_urls(qtbot) -> None:
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    mime = QMimeData()
    mime.setText("du texte, pas des fichiers")
    event = QDragEnterEvent(
        QPoint(5, 5),
        Qt.DropAction.CopyAction,
        mime,
        Qt.MouseButton.LeftButton,
        Qt.KeyboardModifier.NoModifier,
    )
    view.dragEnterEvent(event)
    assert event.isAccepted() is False


def test_tree_widget_emits_dropped_files(qtbot, tmp_path) -> None:
    paths = _make_files(tmp_path, ["tree1.md", "tree2.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    dropped: list[list[str]] = []
    view.tree_explorer.filesDropped.connect(dropped.append)

    enter, _enter_mime = _drag_enter_event(paths)
    view.tree_explorer.dragEnterEvent(enter)
    assert enter.isAccepted()
    drop, _drop_mime = _drop_event(paths)
    view.tree_explorer.dropEvent(drop)

    assert dropped == [paths]


def test_tree_widget_drop_still_supports_internal_move(qtbot) -> None:
    """Le glisser-déposer interne (déplacement de dossier) reste branché."""
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    root = QTreeWidgetItem(["Racine"])
    root.setData(0, Qt.ItemDataRole.UserRole, {"type": "folder", "id": 1})
    child = QTreeWidgetItem(["Enfant"])
    child.setData(0, Qt.ItemDataRole.UserRole, {"type": "doc", "id": 2})
    root.addChild(child)
    view.tree_explorer.addTopLevelItem(root)
    view.tree_explorer.setCurrentItem(child)

    moved: list = []
    view.tree_explorer.filesDropped.connect(lambda p: moved.append(p))
    view.tree_explorer.itemMoved.connect(lambda s, t: moved.append((s, t)))

    mime = QMimeData()
    mime.setText("application/x-ankiforge-internal")
    event = QDropEvent(
        QPointF(1, 1),
        Qt.DropAction.MoveAction,
        mime,
        Qt.MouseButton.LeftButton,
        Qt.KeyboardModifier.NoModifier,
    )
    view.tree_explorer.dropEvent(event)

    assert moved == [({"type": "doc", "id": 2}, {"type": "folder", "id": 1})]


def test_drop_passes_raw_paths_to_the_batch_planner(qtbot, tmp_path) -> None:
    """Les dossiers glissés sont confiés au planificateur, qui les écarte."""
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    with patch.object(view, "_start_batch_import") as start:
        drop, _drop_mime = _drop_event([str(tmp_path)])
        view.dropEvent(drop)

    start.assert_called_once_with([str(tmp_path)])


# ── Déroulement du lot ──────────────────────────────────────────────────────


def test_batch_start_builds_the_queue_and_locks_the_buttons(qtbot, tmp_path) -> None:
    paths = _make_files(tmp_path, ["un.md", "deux.md", "trois.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    with patch(WORKER_PATH) as worker_cls:
        view._start_batch_import([paths[0], paths[1], paths[0]])

    tasks = worker_cls.call_args.args[0]
    assert [t.path for t in tasks] == paths[:2]
    assert [t.index for t in tasks] == [0, 1]
    assert view.btn_import.isEnabled() is False
    assert worker_cls.return_value.start.called is True


def test_batch_start_ignores_missing_files_and_warns(qtbot, tmp_path) -> None:
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    manager = ToastManager.get_instance()
    manager.clear()

    with patch(WORKER_PATH) as worker_cls:
        view._start_batch_import([str(tmp_path / "fantome.md")])

    assert worker_cls.called is False
    assert manager._active_toasts
    manager.clear()


def test_batch_start_refuses_to_run_twice(qtbot, tmp_path) -> None:
    paths = _make_files(tmp_path, ["un.md", "deux.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    _attach_batch_worker(view, paths)

    with patch(WORKER_PATH) as worker_cls:
        view._start_batch_import(paths)

    assert worker_cls.called is False


def test_batch_import_persists_every_extracted_document(qtbot, tmp_path) -> None:
    paths = _make_files(tmp_path, ["lot1.md", "lot2.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    worker.document_started.emit(0, paths[0])
    worker.document_finished.emit(0, paths[0], "lot1", "# lot1\n\nContenu.")
    worker.document_started.emit(1, paths[1])
    worker.document_finished.emit(1, paths[1], "lot2", "# lot2\n\nContenu.")
    worker.batch_finished.emit(2, 0)

    assert DocumentModel.select().where(DocumentModel.title == "lot1").exists()
    assert DocumentModel.select().where(DocumentModel.title == "lot2").exists()
    assert view.btn_import.isEnabled()


def test_batch_import_continues_after_a_failed_document(qtbot, tmp_path) -> None:
    paths = _make_files(tmp_path, ["bon.md", "casse.md", "aussi_bon.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    manager = ToastManager.get_instance()
    manager.clear()

    worker.document_finished.emit(0, paths[0], "bon", "# bon\n\nContenu.")
    worker.document_failed.emit(1, paths[1], "Format de fichier non supporté : .casse")
    worker.document_finished.emit(2, paths[2], "aussi_bon", "# aussi bon\n\nContenu.")
    worker.batch_finished.emit(2, 1)

    assert DocumentModel.select().where(DocumentModel.title == "bon").exists()
    assert DocumentModel.select().where(DocumentModel.title == "aussi_bon").exists()
    assert DocumentModel.select().where(DocumentModel.title == "casse").exists() is False
    assert view.btn_import.isEnabled()
    assert "casse.md" in view.terminal_view.toPlainText()
    manager.clear()


def test_batch_import_progress_is_reported_per_document(qtbot, tmp_path) -> None:
    paths = _make_files(tmp_path, ["un.md", "deux.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    worker.document_started.emit(0, paths[0])
    worker.document_started.emit(1, paths[1])

    terminal = view.terminal_view.toPlainText()
    assert "Importation du document 1 sur 2 : un.md" in terminal
    assert "Importation du document 2 sur 2 : deux.md" in terminal


def test_batch_import_selects_first_imported_document(qtbot, tmp_path) -> None:
    paths = _make_files(tmp_path, ["premier.md", "second.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    worker.document_finished.emit(0, paths[0], "premier", "# premier\n\nContenu.")
    worker.document_finished.emit(1, paths[1], "second", "# second\n\nContenu.")
    worker.batch_finished.emit(2, 0)

    first = DocumentModel.get(DocumentModel.title == "premier")
    assert view._current_doc_id == first.id

    selected_ids = []
    for item in view.tree_explorer.selectedItems():
        data = item.data(0, Qt.ItemDataRole.UserRole)
        if data and data.get("type") == "doc":
            selected_ids.append(data["id"])
    assert selected_ids == [first.id]


def test_batch_summary_toast_reports_ingested_count(qtbot, tmp_path) -> None:
    paths = _make_files(tmp_path, ["un.md", "deux.md", "trois.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    manager = ToastManager.get_instance()
    manager.clear()

    for index, path in enumerate(paths):
        worker.document_finished.emit(index, path, pathlib.Path(path).stem, "# contenu")
    worker.batch_finished.emit(3, 0)

    assert manager._active_toasts
    assert manager._active_toasts[-1].level == ToastLevel.SUCCESS
    assert "3 documents importés !" in view.terminal_view.toPlainText()
    manager.clear()


def test_batch_summary_counts_persist_failures_as_failures(qtbot, tmp_path) -> None:
    """Un document extrait mais non enregistré ne doit pas être compté comme ingéré."""
    paths = _make_files(tmp_path, ["un.md", "deux.md", "trois.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    manager = ToastManager.get_instance()
    manager.clear()

    for index, path in enumerate(paths):
        if index == 1:
            with patch.object(view, "doc_repo") as repo:
                repo.save_imported_document.side_effect = RuntimeError("base verrouillée")
                worker.document_finished.emit(index, path, "deux", "# contenu")
        else:
            worker.document_finished.emit(index, path, pathlib.Path(path).stem, "# contenu")
    worker.batch_finished.emit(3, 0)

    assert manager._active_toasts
    assert manager._active_toasts[-1].level == ToastLevel.WARNING
    assert "2 documents importés, 1 en échec." in view.terminal_view.toPlainText()
    assert DocumentModel.select().where(DocumentModel.title == "deux").exists() is False
    manager.clear()


def test_batch_cancellation_never_reports_a_success_summary(qtbot, tmp_path) -> None:
    """``cancelled`` puis ``batch_finished`` (finally) ne doit pas annoncer un succès."""
    paths = _make_files(tmp_path, ["un.md", "deux.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    manager = ToastManager.get_instance()
    manager.clear()

    worker.document_finished.emit(0, paths[0], "un", "# un")
    worker.cancelled.emit()
    worker.batch_finished.emit(1, 0)

    assert not any(toast.level in (ToastLevel.SUCCESS, ToastLevel.WARNING) for toast in manager._active_toasts)
    assert "Lot terminé" not in view.terminal_view.toPlainText()
    assert "Lot interrompu par l'utilisateur (1 document importé !)" in view.terminal_view.toPlainText()
    manager.clear()


def test_batch_cancellation_reenables_the_import_buttons(qtbot, tmp_path) -> None:
    paths = _make_files(tmp_path, ["un.md", "deux.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    assert view.btn_import.isEnabled() is False
    worker.cancelled.emit()

    assert view.btn_import.isEnabled()
    assert view.btn_import_url.isEnabled()


def test_batch_import_runs_end_to_end_on_real_files(qtbot, tmp_path) -> None:
    """Parcours complet : un vrai lot est extrait dans le thread puis persisté en base."""
    paths = _make_files(tmp_path, ["e2e1.md", "e2e2.md", "e2e3.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    worker = DocumentBatchWorker([DocumentBatchTask(index=i, path=p) for i, p in enumerate(paths)])
    view._bind_batch_worker(worker, worker.tasks)

    with qtbot.waitSignal(worker.batch_finished, timeout=20000) as blocker:
        worker.start()
    worker.wait(5000)
    qtbot.wait(100)

    assert blocker.args == [3, 0]
    for name in ("e2e1", "e2e2", "e2e3"):
        assert DocumentModel.select().where(DocumentModel.title == name).exists()
    assert view.btn_import.isEnabled()
    first = DocumentModel.get(DocumentModel.title == "e2e1")
    assert view._current_doc_id == first.id


def test_batch_import_handles_mixed_formats(qtbot, tmp_path) -> None:
    """Markdown, Word et format inconnu dans le même lot : seuls les fichiers lisibles sont ingérés."""
    import docx

    markdown = tmp_path / "notes.md"
    markdown.write_text("# Notes\n\nContenu markdown.", encoding="utf-8")

    word_path = tmp_path / "cours.docx"
    document = docx.Document()
    document.add_heading("Le théorème", level=1)
    document.add_paragraph("Un triangle rectangle.")
    document.save(str(word_path))

    unreadable = tmp_path / "mystere.inconnu"
    unreadable.write_text("binaire", encoding="utf-8")

    paths = [str(markdown), str(word_path), str(unreadable)]
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    worker = DocumentBatchWorker([DocumentBatchTask(index=i, path=p) for i, p in enumerate(paths)])
    view._bind_batch_worker(worker, worker.tasks)

    with qtbot.waitSignal(worker.batch_finished, timeout=20000) as blocker:
        worker.start()
    worker.wait(5000)
    qtbot.wait(100)

    assert blocker.args == [2, 1]
    assert DocumentModel.get(DocumentModel.title == "notes").file_type == "md"
    assert DocumentModel.get(DocumentModel.title == "cours").file_type == "docx"
    assert DocumentModel.select().where(DocumentModel.title == "mystere").exists() is False
    assert "mystere.inconnu" in view.terminal_view.toPlainText()
    assert "2 documents importés, 1 en échec." in view.terminal_view.toPlainText()


def test_shutdown_cancels_the_running_batch_worker(qtbot, tmp_path) -> None:
    paths = _make_files(tmp_path, ["un.md", "deux.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    started = threading.Event()
    release = threading.Event()

    class _BlockingParser:
        def parse_document(self, source, progress_callback=None, check_cancel=None):
            started.set()
            release.wait(5)
            return "# contenu"

    worker = DocumentBatchWorker([DocumentBatchTask(index=i, path=p) for i, p in enumerate(paths)])
    view._bind_batch_worker(worker, worker.tasks)
    try:
        with patch("ankiforge.services.workers.document_batch_worker.DocumentParser", return_value=_BlockingParser()):
            worker.start()
            assert started.wait(5)
            view.shutdown()
            assert worker.is_cancelled() is True
    finally:
        release.set()
        worker.wait(5000)


def test_batch_import_uses_unique_titles_when_duplicated(qtbot, tmp_path) -> None:
    """Deux fichiers homonymes produisent deux fiches distinctes (contrainte d'unicité)."""
    first_dir = tmp_path / "un"
    second_dir = tmp_path / "deux"
    first_dir.mkdir()
    second_dir.mkdir()
    paths = [str(first_dir / "memo.md"), str(second_dir / "memo.md")]
    for path in paths:
        with open(path, "w", encoding="utf-8") as handle:
            handle.write("# memo\n\nContenu.")

    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    worker.document_finished.emit(0, paths[0], "memo", "# memo\n\nContenu.")
    worker.document_finished.emit(1, paths[1], "memo", "# memo\n\nContenu.")
    worker.batch_finished.emit(2, 0)

    titles = [doc.title for doc in DocumentModel.select().where(DocumentModel.title.startswith("memo"))]
    assert len(titles) == 2
    assert titles[0] != titles[1]


# ── Message de récapitulatif ────────────────────────────────────────────────


@pytest.mark.parametrize(
    ("ok", "failed", "expected"),
    [
        (1, 0, "1 document importé !"),
        (3, 0, "3 documents importés !"),
        (2, 1, "2 documents importés, 1 en échec."),
    ],
)
def test_batch_summary_message(ok: int, failed: int, expected: str) -> None:
    assert batch_summary_message(ok, failed) == expected


# ── Tests TDD : Rafraîchissement incrémental, Fiche PDF provisoire & Progression ─────────


def test_batch_document_finished_refreshes_tree_incrementally_preserving_selection(qtbot, tmp_path) -> None:
    """Chaque document terminé apparaît immédiatement dans l'arbre sans attendre la fin du lot,

    et la sélection reste fixée sur le premier document ingéré.
    """
    paths = _make_files(tmp_path, ["premier.md", "second.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    # Document 1 démarre et finit
    worker.document_started.emit(0, paths[0])
    worker.document_finished.emit(0, paths[0], "premier", "# premier\n\nContenu 1.")

    # L'arbre doit DÉJÀ contenir le premier document AVANT que le lot soit fini
    first_doc = DocumentModel.get(DocumentModel.title == "premier")
    assert view.tree_explorer.topLevelItemCount() == 1
    assert "premier" in view.tree_explorer.topLevelItem(0).text(0)
    assert view._current_doc_id == first_doc.id

    # Document 2 démarre et finit
    worker.document_started.emit(1, paths[1])
    worker.document_finished.emit(1, paths[1], "second", "# second\n\nContenu 2.")

    # L'arbre contient maintenant les 2 documents
    assert view.tree_explorer.topLevelItemCount() == 2
    # La sélection et _current_doc_id restent bien sur le 1er document
    assert view._current_doc_id == first_doc.id
    selected_items = view.tree_explorer.selectedItems()
    assert len(selected_items) == 1
    assert selected_items[0].data(0, Qt.ItemDataRole.UserRole)["id"] == first_doc.id

    worker.batch_finished.emit(2, 0)
    assert view._current_doc_id == first_doc.id


def test_batch_pdf_creates_provisional_doc_at_started_and_updates_at_finished(qtbot, tmp_path) -> None:
    """Un PDF du lot dispose d'une fiche provisoire dès document_started (analyse en cours),

    puis est mis à jour à document_finished avec son contenu extrait et ses chunks.
    """
    pdf_file = tmp_path / "cours_bio.pdf"
    pdf_file.write_bytes(b"%PDF-1.4 Minimal PDF dummy content for test")
    paths = [str(pdf_file)]

    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    # Démarrage de l'extraction du PDF
    worker.document_started.emit(0, str(pdf_file))

    # Une fiche provisoire doit être créée en base dès le démarrage
    assert DocumentModel.select().where(DocumentModel.title == "cours_bio").exists()
    prov_doc = DocumentModel.get(DocumentModel.title == "cours_bio")
    assert prov_doc.file_type == "pdf"
    assert prov_doc.content == ""
    assert prov_doc.original_media is not None

    # L'arbre doit refléter la fiche avec l'état 'analyse en cours'
    assert view.tree_explorer.topLevelItemCount() == 1
    item = view.tree_explorer.topLevelItem(0)
    assert "(analyse en cours)" in item.text(0)

    # Fin de l'extraction : le contenu extrait arrive
    extracted_content = "# Cours de Biologie\n\nContenu extrait via Marker OCR."
    worker.document_finished.emit(0, str(pdf_file), "cours_bio", extracted_content)

    # La fiche existante a été mise à jour (même ID, pas de doublon)
    updated_doc = DocumentModel.get(DocumentModel.title == "cours_bio")
    assert updated_doc.id == prov_doc.id
    assert updated_doc.content == extracted_content
    # L'arbre ne porte plus la mention 'analyse en cours' (nouvel item suite au refresh)
    updated_item = view.tree_explorer.topLevelItem(0)
    assert "(analyse en cours)" not in updated_item.text(0)

    worker.batch_finished.emit(1, 0)


def test_batch_progress_indicator_displays_x_sur_n_and_toggles_visibility(qtbot, tmp_path) -> None:
    """Un indicateur de progression X sur N est visible dans la vue pendant le lot et masqué à la fin."""
    paths = _make_files(tmp_path, ["doc1.md", "doc2.md", "doc3.md"])
    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    view.show()

    # Au départ, le conteneur de progression est masqué
    assert view.batch_progress_container.isVisible() is False

    worker = _attach_batch_worker(view, paths)
    assert view.batch_progress_container.isVisible() is True
    assert view.batch_progress_bar.maximum() == 3

    # Document 1 en cours
    worker.document_started.emit(0, paths[0])
    assert "1 sur 3" in view.lbl_batch_progress.text()
    assert "1 / 3" in view.badge_batch_count.text()
    assert view.batch_progress_bar.value() == 1

    # Document 2 en cours
    worker.document_started.emit(1, paths[1])
    assert "2 sur 3" in view.lbl_batch_progress.text()
    assert "2 / 3" in view.badge_batch_count.text()
    assert view.batch_progress_bar.value() == 2

    # Clôture du lot : le conteneur est masqué
    worker.batch_finished.emit(2, 0)
    assert view.batch_progress_container.isVisible() is False


def test_batch_pdf_failure_transitions_to_not_extracted(qtbot, tmp_path) -> None:
    """Si l'extraction d'un PDF échoue, la fiche provisoire passe en état (Non extrait)."""
    pdf_file = tmp_path / "corrompu.pdf"
    pdf_file.write_bytes(b"%PDF-corrompu")
    paths = [str(pdf_file)]

    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)
    worker = _attach_batch_worker(view, paths)

    worker.document_started.emit(0, str(pdf_file))
    assert DocumentModel.select().where(DocumentModel.title == "corrompu").exists()

    # Échec du document
    worker.document_failed.emit(0, str(pdf_file), "Erreur de décodage PDF Marker")

    item = view.tree_explorer.topLevelItem(0)
    assert "(Non extrait)" in item.text(0)
    assert "(analyse en cours)" not in item.text(0)

    worker.batch_finished.emit(0, 1)


def test_single_pdf_import_uses_persist_imported_document_and_preserves_flow(qtbot, tmp_path) -> None:
    """L'import unitaire d'un PDF utilise le point d'écriture unifié _persist_imported_document."""
    pdf_file = tmp_path / "cours_unitaire.pdf"
    pdf_file.write_bytes(b"%PDF-1.4 dummy content")

    view = DocumentsView(ai_manager=None)
    qtbot.addWidget(view)

    with patch.object(view, "_start_document_worker") as start_worker:
        view._import_pdf_directly(str(pdf_file))

    assert DocumentModel.select().where(DocumentModel.title == "cours_unitaire").exists()
    doc = DocumentModel.get(DocumentModel.title == "cours_unitaire")
    assert doc.file_type == "pdf"
    assert doc.content == ""
    assert start_worker.called
    assert start_worker.call_args.kwargs.get("doc_id") == doc.id
